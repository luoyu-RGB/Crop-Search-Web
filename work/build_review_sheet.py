"""把 data/extracted/page-NN.json 汇总为人工校对用的 Excel 工作簿与 CSV。

输出：
  data/review/作物资料校对表.xlsx   —— 5 张工作表：作物主表 / 别名 / 采集点 / 含油率与化学组分 / 图片
  data/review/crops.csv             —— 作物主表（便于版本对比）

用法：python work/build_review_sheet.py
"""
import glob
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT_DIR = os.path.join('data', 'review')
os.makedirs(OUT_DIR, exist_ok=True)

CROP_HEADERS = [
    ('crop_no', 8), ('page', 6), ('printed_page', 8), ('chinese_name', 14), ('alias_text', 14),
    ('english_name', 18), ('latin_name', 30), ('family', 12), ('family_latin', 16), ('genus', 12),
    ('features', 60), ('distribution', 60), ('cultivation', 60), ('uses', 40),
    ('image_count', 10), ('chem_rows', 10), ('uncertain_fields', 24),
]
ALIAS_HEADERS = [('crop_no', 8), ('chinese_name', 14), ('alias_name', 18), ('alias_type', 12), ('remark', 20)]
POINT_HEADERS = [
    ('crop_no', 8), ('chinese_name', 14), ('province', 10), ('locality', 30), ('latitude', 14),
    ('longitude', 14), ('altitude', 12), ('date', 12), ('collectors', 24), ('collection_no', 16),
]
CHEM_HEADERS = [
    ('crop_no', 8), ('chinese_name', 14), ('page', 6), ('collection_unit', 14), ('test_unit', 12),
    ('test_part', 10), ('origin', 14), ('oil_content', 12), ('iodine_value', 12), ('acid_value', 12),
    ('saponification_value', 16),
    ('C12:0', 9), ('C14:0', 9), ('C16:0', 9), ('C16:1', 9), ('C18:0', 9),
    ('C18:1', 9), ('C18:2', 9), ('C18:3', 9), ('C20:0', 9), ('C20:1', 9),
]
IMAGE_HEADERS = [
    ('crop_no', 8), ('chinese_name', 14), ('page', 6), ('sort_order', 10),
    ('image_file', 34), ('image_title', 24), ('image_type', 12), ('is_cover', 9),
]

FA_KEYS = ['C12:0', 'C14:0', 'C16:0', 'C16:1', 'C18:0', 'C18:1', 'C18:2', 'C18:3', 'C20:0', 'C20:1']


def style_sheet(ws, headers):
    fill = PatternFill('solid', fgColor='1F5C3A')
    for col, (title, width) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = 'A2'


def guess_alias_type(alias):
    if any('\u4e00' <= ch <= '\u9fff' for ch in alias):
        return '中文别名'
    return '外文别名'


wb = Workbook()

ws_crop = wb.active
ws_crop.title = '作物主表'
style_sheet(ws_crop, CROP_HEADERS)
ws_alias = wb.create_sheet('别名')
style_sheet(ws_alias, ALIAS_HEADERS)
ws_point = wb.create_sheet('采集点')
style_sheet(ws_point, POINT_HEADERS)
ws_chem = wb.create_sheet('含油率与化学组分')
style_sheet(ws_chem, CHEM_HEADERS)
ws_image = wb.create_sheet('图片')
style_sheet(ws_image, IMAGE_HEADERS)

rows_csv = []
crop_no = 0
total_chem = 0

for path in sorted(glob.glob(os.path.join('data', 'extracted', 'page-*.json'))):
    with open(path, encoding='utf-8') as fh:
        doc = json.load(fh)
    page = doc.get('page')
    printed = doc.get('printedPage')
    for crop in doc.get('crops', []):
        crop_no += 1
        alias_text = crop.get('aliasText') or ''
        chem = crop.get('chemistry') or []
        images = crop.get('images') or []
        uncertain = crop.get('uncertain') or []

        ws_crop.append([
            crop_no, page, printed, crop.get('chineseName'), alias_text,
            crop.get('englishName'), crop.get('latinName'), crop.get('family'),
            crop.get('familyLatin'), crop.get('genus'),
            crop.get('features'), crop.get('distribution'), crop.get('cultivation'), crop.get('uses'),
            len(images), len(chem), '、'.join(uncertain) if isinstance(uncertain, list) else uncertain,
        ])
        rows_csv.append({
            'crop_no': crop_no, 'page': page, 'chinese_name': crop.get('chineseName'),
            'latin_name': crop.get('latinName'), 'family': crop.get('family'), 'genus': crop.get('genus'),
        })

        for alias in [a.strip() for a in alias_text.replace('，', '、').split('、') if a.strip()]:
            ws_alias.append([crop_no, crop.get('chineseName'), alias, guess_alias_type(alias), ''])

        for pt in crop.get('distributionPoints') or []:
            ws_point.append([
                crop_no, crop.get('chineseName'), pt.get('province'), pt.get('locality'),
                pt.get('latitude'), pt.get('longitude'), pt.get('altitude'), pt.get('date'),
                pt.get('collectors'), pt.get('collectionNo'),
            ])

        for row in chem:
            fa = row.get('fattyAcids') or {}
            ws_chem.append([
                crop_no, crop.get('chineseName'), page,
                row.get('collectionUnit'), row.get('testUnit'), row.get('testPart'), row.get('origin'),
                row.get('oilContent'), row.get('iodineValue'), row.get('acidValue'),
                row.get('saponificationValue'),
            ] + [fa.get(k) for k in FA_KEYS])
            total_chem += 1

        for img in images:
            ws_image.append([
                crop_no, crop.get('chineseName'), page, img.get('order'),
                img.get('file'), img.get('description'),
                img.get('type'), '是' if img.get('isCover') else '否',
            ])

for ws, headers in ((ws_crop, CROP_HEADERS), (ws_chem, CHEM_HEADERS), (ws_point, POINT_HEADERS),
                    (ws_image, IMAGE_HEADERS), (ws_alias, ALIAS_HEADERS)):
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=cell.column in (11, 12, 13) if ws is ws_crop else False)

xlsx_path = os.path.join(OUT_DIR, '作物资料校对表.xlsx')
wb.save(xlsx_path)

csv_path = os.path.join(OUT_DIR, 'crops.csv')
with open(csv_path, 'w', encoding='utf-8-sig', newline='') as fh:
    import csv
    writer = csv.DictWriter(fh, fieldnames=list(rows_csv[0].keys()) if rows_csv else ['crop_no'])
    writer.writeheader()
    writer.writerows(rows_csv)

print(f'作物条目: {crop_no}')
print(f'化学数据行: {total_chem}')
print(f'已生成: {xlsx_path}')
print(f'已生成: {csv_path}')
