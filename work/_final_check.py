"""统计将纳入 Git 的文件规模，并检查校对表内容。"""
import os
import subprocess

from openpyxl import load_workbook

out = subprocess.run(['git', 'status', '--porcelain', '-uall'], capture_output=True, text=True, encoding='utf-8')
files = []
for line in out.stdout.splitlines():
    p = line[3:].strip().strip('"')
    if os.path.isfile(p):
        files.append(p)
total = sum(os.path.getsize(p) for p in files)
print(f'将纳入 Git 的文件数: {len(files)}，合计 {total / 1024 / 1024:.2f} MB')
big = sorted(((os.path.getsize(p), p) for p in files), reverse=True)[:5]
for size, p in big:
    print(f'  {size / 1024 / 1024:6.2f} MB  {p}')

wb = load_workbook('data/review/作物资料校对表.xlsx')
print('\n工作表:', wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print(f'  {name}: {ws.max_row - 1} 行 x {ws.max_column} 列')
ws = wb['作物主表']
print('\n作物主表前 3 行:')
for row in ws.iter_rows(min_row=1, max_row=4, max_col=8, values_only=True):
    print('  ', row)
